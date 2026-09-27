"""Rebuild the annotated DOCX and Markdown from the adjacent evidence register.
Usage: bundled_python build_bibliography.py PATH_TO_HANDOFF_FOLDER
"""
from pathlib import Path
import json, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
folder=Path(sys.argv[1])
data=json.loads((folder/'evidence-register.json').read_text())
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(5)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 2'].font.size=Pt(13)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('Zombie Fire Sim research bibliography  •  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title='Zombie Fire Sim annotated physics bibliography'
doc.core_properties.subject='Source audit and research starting points for the coupled simulator'
doc.core_properties.author='Codex'
md=[]
def heading(text,level=1):
    doc.add_heading(text,level);md.append('#'*(level+1)+' '+text+'\n')
def para(text,label=None):
    p=doc.add_paragraph()
    if label:p.add_run(label+': ').bold=True
    p.add_run(text)
    md.append(('**'+label+'** ' if label else '')+text+'\n')
    return p
def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1F4E79');rp.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

doc.add_paragraph('Zombie Fire Sim annotated physics bibliography','Title');md.append('# Zombie Fire Sim annotated physics bibliography\n')
para('Research sources and evidence gaps for the next coupled implementation')
para('Prepared 27 September 2026 for Kaden Smith. This bibliography maps 31 peer-reviewed papers to the current simulator or proposed extensions. It includes selected published observations with their conditions and source locations. It supplies starting evidence for implementation and testing; it does not establish experimental validation of the complete scene.')
heading('How to use this bibliography')
para('Each entry states the source, relevant finding, code use, limitations and depth of access. “Current” means a citation or method already appears in the project; it does not mean all parameters or claims have been verified. “Proposed” means no implementation or validation is claimed. Abstract-only entries cannot support extraction of unseen coefficients. Reviews and fitted model values are distinguished from original measurements.')
para('Construct separate literature benchmark cases before combining physics. Preserve peat type, preparation, density, moisture basis, geometry, oxygen access, thermal conditions and measurement method. Combine studies only when their conditions support it. Retain paired observations and distinguish calibration from independent holdout tests.')
heading('Requested research scenarios')
para('The user confirmed three terrain cases: rooted peat, layered peat/mineral soil, and soil with a rocky subsurface. Hole depth, diameter and water inlet temperature must be adjustable. Proposed water input is 100 US gallons/min (6.309 L/s), using a provisional US-gallon interpretation. Flow duration, thermal state and delivery capability remain scenario inputs without measurements. A 0.5 m sphere diameter is a provisional interpretation of the requested source size.')
para('Use controlled non-explosive heating and a vented or pressure-relieved research configuration. Report per-output uncertainty after each run as computed percentage ranges where meaningful, with an absolute-range fallback near zero. Unsupported model behavior and numerical failure must remain visible independently of those percentages.')
heading('Evidence status')
para('The search was targeted, not exhaustive. No original machine-readable raw experimental dataset was downloaded. Some full papers were inaccessible. The adjacent evidence-register.json preserves detailed source and access records; RESEARCH-SEARCH-LOG.md states coverage and the required next search. The final section lists technical references outside the peer-reviewed paper set and unresolved property audits.')

short={'D25':'Soil organic matter mixture theory','AT23':'Thermal diffusivity of peat sand mixtures','P17':'Field experiments in organic soils','OD09':'Moisture and organic soil thermal conductivity','MA07':'Engineering properties of fibrous peat','M22':'Laboratory water injection in peat','DB25':'Peat type and combustion moisture thresholds'}
for i,e in enumerate(data['records']):
    if i%2==0:doc.add_page_break()
    heading(f"{i+1} {short.get(e['id'],e['title'])}")
    para(e['citation'])
    p=doc.add_paragraph();link(p,'DOI '+e['doi'],e['url'])
    md.append('[DOI '+e['doi']+']('+e['url']+')\n')
    urls=[]
    for u in [e.get('fulltext_url'),e.get('institutional_url')]+e.get('additional_links',[]):
        if u and u!=e['url'] and u not in urls:urls.append(u)
    for n,u in enumerate(urls):
        p.add_run('  |  ');link(p,'Source record' if n==0 else 'Linked repository',u)
        md.append('[Source record]('+u+')\n')
    para(e['role'],'Role')
    para(e['summary'],'Finding')
    para(e['use'],'Use in the simulator')
    for ob in e.get('observations',[]):
        vals='; '.join(k.replace('_',' ')+' = '+str(v) for k,v in ob['values'].items())
        para(ob.get('display_summary',ob['location']+'. '+ob['basis']+'. '+vals+'.'),'Inspected observations')
    para(e['limits'],'Limits')
    para(e['access'],'Access')
    p=para('; '.join(e['code_locations']),'Code locations')
    for r in p.runs:r.font.size=Pt(9)

doc.add_page_break();heading('Technical references requiring a separate evidence audit')
para('The current coupled documentation also cites the following technical sources. They are useful for formulation or numerical verification, but they are not peer-reviewed empirical measurements. The next implementation must trace physical data and claims to applicable peer-reviewed papers and verify any adopted correlations.')
for e in data['non_peer_reviewed_current_references']:
    para(e['citation']);p=doc.add_paragraph();link(p,'Reference link',e['url']);md.append('[Reference link]('+e['url']+')\n')
    para(e['use']+' '+e['limit'])
heading('Unresolved properties and validation')
for t in data['unresolved_property_audit']:
    p=doc.add_paragraph(t,style='List Bullet');md.append('- '+t+'\n')
para('The current spatial fracture mesh-energy check fails its stated gate, and a default coupled fracture attempt fails an energy increment check. These are numerical failures in addition to material evidence gaps. Keep the capability gated until the stated checks pass; attractive animation or an AI score cannot replace them.')
(folder/'BIBLIOGRAPHY.md').write_text('\n'.join(md))
doc.save(folder/'Annotated-Physics-Bibliography.docx')
print('Created bibliography for',len(data['records']),'papers')
